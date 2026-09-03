import json
import os
import re
from io import BytesIO

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document

MODEL = "gemini-3.5-flash"

st.set_page_config(page_title="AI Resume ATS Analyzer", page_icon="📄", layout="wide")

st.markdown("""
<style>
.main-title {font-size: 2.6rem; font-weight: 800; margin-bottom: .2rem;}
.subtitle {color: #6b7280; font-size: 1.05rem; margin-bottom: 1.5rem;}
.score {font-size: 4rem; font-weight: 800; line-height: 1;}
.card {padding: 1.1rem 1.2rem; border: 1px solid rgba(128,128,128,.25); border-radius: 14px;}
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">📄 AI Resume ATS Analyzer</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="subtitle">Upload your resume and get an AI-powered ATS score, weaknesses, and actionable improvements.</div>',
    unsafe_allow_html=True,
)

def get_api_key():
    try:
        return st.secrets["GEMINI_API_KEY"]
    except Exception:
        return os.getenv("GEMINI_API_KEY", "")

def extract_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()

def extract_docx(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip()

def extract_resume(uploaded_file) -> str:
    data = uploaded_file.getvalue()
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    raise ValueError("Unsupported file type.")

def normalize_result(result):
    # Gemini should return JSON, but this makes the app resilient to fenced JSON.
    text = result.text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text)
    return json.loads(text)

def analyze_resume(resume_text: str, job_description: str):
    client = genai.Client(api_key=get_api_key())

    prompt = f"""
You are an expert ATS resume evaluator and technical recruiter.

Analyze the resume below. If a job description is provided, score the resume
against that specific job; otherwise score it against general ATS best practices.

Important:
- The ATS score is a practical heuristic, NOT a score from a real ATS vendor.
- Do not invent experience, education, skills, dates, metrics, employers, or achievements.
- Identify missing keywords only when they are genuinely relevant to the supplied job description.
- Evaluate parsing/readability, section structure, contact information, skills,
  experience bullets, measurable impact, keyword alignment, consistency, and formatting.
- Give specific, actionable improvements.
- Return ONLY valid JSON matching the schema below.

JSON schema:
{{
  "ats_score": 0,
  "score_label": "Poor|Needs Improvement|Good|Excellent",
  "summary": "short overall assessment",
  "breakdown": {{
    "keyword_match": 0,
    "format_parsability": 0,
    "experience_impact": 0,
    "skills_relevance": 0,
    "section_completeness": 0
  }},
  "strengths": ["..."],
  "improvements": [
    {{
      "priority": "High|Medium|Low",
      "issue": "...",
      "recommendation": "...",
      "example": "..."
    }}
  ],
  "missing_keywords": ["..."],
  "ats_checklist": {{
    "standard_headings": true,
    "contact_information": true,
    "readable_structure": true,
    "keyword_alignment": true,
    "quantified_achievements": true,
    "consistent_dates": true
  }}
}}

Resume:
---BEGIN RESUME---
{resume_text}
---END RESUME---

Job description (may be empty):
---BEGIN JOB DESCRIPTION---
{job_description}
---END JOB DESCRIPTION---
"""

    response = client.models.generate_content(
        model=MODEL,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    result = normalize_result(response)

    # Keep the score trustworthy and bounded.
    score = int(result.get("ats_score", 0))
    result["ats_score"] = max(0, min(100, score))
    return result

uploaded = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="PDF or DOCX. For best ATS analysis, use a text-based PDF rather than a scanned image.",
)

job_description = st.text_area(
    "Job description (optional)",
    placeholder="Paste the target job description here for a more targeted ATS score...",
    height=180,
)

if uploaded:
    st.caption(f"Selected: {uploaded.name}")

if st.button("🔍 Analyze Resume", type="primary", use_container_width=True):
    if not uploaded:
        st.warning("Please upload a PDF or DOCX resume first.")
        st.stop()

    if not get_api_key():
        st.error("Gemini API key is missing. Add GEMINI_API_KEY to Streamlit Secrets or your environment variables.")
        st.stop()

    try:
        with st.spinner("Reading and analyzing your resume with Gemini Flash..."):
            resume_text = extract_resume(uploaded)

            if len(resume_text.strip()) < 80:
                st.error("Very little text could be extracted. If this is a scanned PDF, upload a text-based PDF or DOCX.")
                st.stop()

            result = analyze_resume(resume_text, job_description.strip())

        st.success("Analysis complete.")

        score = result["ats_score"]
        col1, col2 = st.columns([1, 3])
        with col1:
            st.markdown('<div class="card">', unsafe_allow_html=True)
            st.markdown("### ATS Score")
            st.markdown(f'<div class="score">{score}<span style="font-size:1.4rem;">/100</span></div>', unsafe_allow_html=True)
            st.markdown(f"**{result.get('score_label', '')}**")
            st.markdown("</div>", unsafe_allow_html=True)
        with col2:
            st.markdown("### Overall assessment")
            st.write(result.get("summary", ""))

        st.divider()

        st.subheader("📊 Score Breakdown")
        breakdown = result.get("breakdown", {})
        cols = st.columns(5)
        labels = [
            ("Keyword Match", "keyword_match"),
            ("Format", "format_parsability"),
            ("Impact", "experience_impact"),
            ("Skills", "skills_relevance"),
            ("Completeness", "section_completeness"),
        ]
        for col, (label, key) in zip(cols, labels):
            value = max(0, min(100, int(breakdown.get(key, 0))))
            col.metric(label, f"{value}/100")

        left, right = st.columns(2)
        with left:
            st.subheader("✅ Strengths")
            for item in result.get("strengths", []):
                st.markdown(f"- {item}")

        with right:
            st.subheader("🔑 Missing / useful keywords")
            keywords = result.get("missing_keywords", [])
            if keywords:
                st.write(", ".join(keywords))
            else:
                st.write("No major missing keywords were identified.")

        st.subheader("🛠️ Recommended Improvements")
        improvements = result.get("improvements", [])
        for i, item in enumerate(improvements, 1):
            priority = item.get("priority", "Medium")
            with st.expander(f"{i}. [{priority}] {item.get('issue', 'Improvement')}"):
                st.write(item.get("recommendation", ""))
                if item.get("example"):
                    st.markdown("**Example:**")
                    st.info(item["example"])

        st.subheader("☑️ ATS Checklist")
        checklist = result.get("ats_checklist", {})
        checklist_labels = {
            "standard_headings": "Standard section headings",
            "contact_information": "Contact information present",
            "readable_structure": "Readable ATS-friendly structure",
            "keyword_alignment": "Keyword alignment",
            "quantified_achievements": "Quantified achievements",
            "consistent_dates": "Consistent dates",
        }
        for key, label in checklist_labels.items():
            st.checkbox(label, value=bool(checklist.get(key, False)), disabled=True)

        st.caption("Note: This score is an AI-assisted heuristic. Different ATS platforms and employers use different parsing and ranking rules.")

    except json.JSONDecodeError:
        st.error("Gemini returned an invalid analysis format. Please try again.")
    except Exception as exc:
        st.error(f"Analysis failed: {exc}")
else:
    st.info("Upload a PDF or DOCX resume above to begin.")
